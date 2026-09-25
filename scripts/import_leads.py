#!/usr/bin/env python3
"""从2026年1-9月线索底表_合并版.xlsx导入线索到CRM数据库（修正版）"""
import openpyxl
import sqlite3
from datetime import datetime

DB = '/home/mlclaw/.openclaw/workspace/mkt-crm/data/mkt-crm.db'
XLSX = '/home/mlclaw/.openclaw/workspace/data/2026年1-9月线索底表_合并版.xlsx'

def gv(r, *keys, default=''):
    """从dict中按多个候选key取值"""
    for k in keys:
        v = r.get(k)
        if v is not None and str(v).strip() != '':
            return v
    return default

def norm_date(v):
    if not v: return ''
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    if len(s) >= 10:
        for fmt in ('%Y/%m/%d', '%Y-%m-%d', '%Y.%m.%d'):
            try: return datetime.strptime(s[:10], fmt).strftime('%Y-%m-%d')
            except: pass
    return s[:10]

def norm_phone(v):
    if v is None: return ''
    if isinstance(v, float):
        if v.is_integer(): return str(int(v))
        return str(v)
    if isinstance(v, int): return str(v)
    s = str(v).strip()
    if s.endswith('.0'): s = s[:-2]
    return s

def classify(r):
    """根据字段判断线索状态"""
    valid = str(gv(r, '有效性', '是否有效', default='')).strip()
    if valid == '无效': return 'closed_lost'
    transfer = gv(r, '转出日期')
    if transfer and str(transfer).strip(): return 'qualified'
    opp = gv(r, '商机号')
    if opp and str(opp).strip(): return 'opportunity'
    assigned = gv(r, '分配销售')
    if assigned and str(assigned).strip():
        if valid in ('是',) or valid == '': return 'contacted'
    amount = gv(r, '成单金额')
    try:
        if amount and float(str(amount).replace(',', '')) > 0: return 'closed_won'
    except: pass
    if valid == '待确认' or not valid:
        # 没有分配也没有转出 = 新进线
        return 'new'
    return 'new'

def extract_row(r, source_default=''):
    company = str(gv(r, '公司名称', '公司', default='')).strip()
    contact = str(gv(r, '姓名', '联系人', default='')).strip()
    if not company and not contact: return None

    phone = norm_phone(gv(r, '电话', '手机号'))
    amount_raw = gv(r, '成单金额', default=0)
    try: amount = float(str(amount_raw).replace(',', '') or 0)
    except: amount = 0
    source = str(gv(r, '线索来源', '来源', default=source_default)).strip()
    # 来源规范化
    if '晚点' in source: source = 'Octo·晚点头条'
    elif source in ('','400','400来电'): source = '400电话'

    follow = str(gv(r, '跟进情况', '跟进反馈', '沟通内容', default='')).strip()

    return (
        company[:100],
        contact[:50],
        str(gv(r, '职位', default=''))[:50],
        phone,
        source[:50],
        str(gv(r, '了解渠道', default=''))[:100],
        classify(r),
        str(gv(r, '匹配产品', '需求产品', default=''))[:100],
        '',  # industry
        str(gv(r, '所属团队', '团队', '部门', default=''))[:50],
        str(gv(r, '分配销售', default=''))[:50],
        norm_date(gv(r, '进线日期', '咨询日期', '日期', default='')),
        norm_date(gv(r, '转出日期', default='')),
        str(gv(r, '商机号', default=''))[:50],
        str(gv(r, '商机阶段', default=''))[:50],
        amount,
        str(gv(r, '丢单原因', default=''))[:200],
        follow[:2000],
        None
    )

conn = sqlite3.connect(DB)
cur = conn.cursor()
cur.execute("DELETE FROM leads")
print("清空旧线索完成")

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
total = 0
stats = {}

for sn, src_default in [
    ('转出线索', None),
    ('400数据', '400电话'),
    ('官微数据', '官方微信'),
    ('octo数据', 'Octo体验'),
    ('Octo&晚点头条数据', None),
    ('其他渠道数据', None),
]:
    if sn not in wb.sheetnames: continue
    ws = wb[sn]
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(c).strip() if c else '' for c in next(rows_iter)]
    cnt = 0
    batch = []
    for row in rows_iter:
        if not any(v is not None and str(v).strip() for v in row): continue
        r = dict(zip(headers, row))
        rec = extract_row(r, src_default)
        if rec is None: continue
        batch.append(rec)
        stats[rec[6]] = stats.get(rec[6], 0) + 1
        cnt += 1
    if batch:
        cur.executemany("""INSERT INTO leads (company_name, contact_name, contact_title, phone, source_channel, source_detail,
            status, product, industry, team, assigned_to, inbound_date, transfer_date,
            opportunity_id, opportunity_stage, deal_amount, lost_reason, requirement, account_id)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", batch)
    total += cnt
    print(f"{sn}: {cnt}条")

wb.close()
conn.commit()

print(f"\n合计导入: {total}条")
print("状态分布:")
names = {'new':'新进线','contacted':'已联系','qualified':'已转出','opportunity':'商机','closed_won':'成单','closed_lost':'无效/丢单'}
for s in ['new','contacted','qualified','opportunity','closed_won','closed_lost']:
    c = stats.get(s, 0)
    if c: print(f"  {names[s]}: {c}条")

# 来源分布
print("\n来源Top10:")
cur.execute("SELECT source_channel, COUNT(*) c FROM leads WHERE source_channel != '' GROUP BY source_channel ORDER BY c DESC LIMIT 10")
for src, c in cur.fetchall():
    print(f"  {src}: {c}条")

conn.close()
